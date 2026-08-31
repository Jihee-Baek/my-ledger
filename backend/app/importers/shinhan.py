"""신한카드 '카드이용내역조회' Excel export (real .xlsx, single flat sheet).

Unlike the older '이용대금명세서' PDF-style export, this format has one row
per transaction with a real 승인번호 (approval number), which is what makes
reliable de-duplication possible for 신한카드. The same column layout
covers both domestic and overseas rows: 거래통화/해외이용금액 are only
populated for overseas transactions, while 금액 is always the KRW amount
billed to the card.
"""

from pathlib import Path

import openpyxl

from app.importers.base import ParsedTransaction
from app.importers.normalize import (
    is_cancelled,
    mask_card_number_if_needed,
    parse_amount,
    parse_date,
)
from app.importers.sniff import is_real_xlsx

_EXPECTED_HEADER = {
    "거래일", "카드구분", "이용카드", "가맹점명", "업종", "승인번호", "금액",
    "매입구분", "이용구분", "거래통화", "최초결제일자", "해외이용금액", "취소상태",
}


class ShinhanCardImporter:
    source = "shinhan_card"

    def can_handle(self, file_path: Path) -> bool:
        if not is_real_xlsx(file_path):
            return False
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True)
            header = next(wb.active.iter_rows(values_only=True))
        except Exception:
            return False
        return _EXPECTED_HEADER.issubset(set(header))

    def parse(self, file_path: Path) -> list[ParsedTransaction]:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        header, *data_rows = rows

        results = []
        for values in data_rows:
            row = dict(zip(header, values))
            if not row.get("거래일"):  # trailing "총 N건" summary row
                continue

            local_amount = row.get("해외이용금액")
            local_currency = row.get("거래통화")
            description = None
            if local_amount not in (None, ""):
                description = f"{row.get('업종') or ''} / 현지금액 {local_amount} {local_currency}".strip()
            elif row.get("업종"):
                description = str(row["업종"])

            results.append(
                ParsedTransaction(
                    transaction_date=parse_date(str(row["거래일"])),
                    transaction_type="EXPENSE",
                    amount=parse_amount(row["금액"]),
                    currency="KRW",
                    merchant_raw=str(row["가맹점명"]).strip(),
                    card_institution="신한카드",
                    card_number_masked=mask_card_number_if_needed(str(row["이용카드"])),
                    source_transaction_id=str(row["승인번호"]),
                    description=description,
                    is_excluded=is_cancelled(row.get("취소상태")),
                    raw_row=row,
                )
            )
        return results
