"""삼성카드 Excel export (real .xlsx). Domestic and overseas downloads share
the same two-sheet shape (a one-row summary sheet + a detail sheet) but the
detail sheet name and columns differ, so both variants are handled here."""

from pathlib import Path

import openpyxl

from app.importers.base import ParsedTransaction
from app.importers.normalize import (
    parse_time,
    is_cancelled,
    mask_card_number_if_needed,
    parse_amount,
    parse_date,
)
from app.importers.sniff import is_real_xlsx

_DOMESTIC_SHEET = "■ 국내이용내역"
_OVERSEAS_SHEET = "■ 해외이용내역"


class SamsungCardImporter:
    source = "samsung_card"

    def can_handle(self, file_path: Path) -> bool:
        if not is_real_xlsx(file_path):
            return False
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True)
        except Exception:
            return False
        return _DOMESTIC_SHEET in wb.sheetnames or _OVERSEAS_SHEET in wb.sheetnames

    def parse(self, file_path: Path) -> list[ParsedTransaction]:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        if _DOMESTIC_SHEET in wb.sheetnames:
            return self._parse_domestic(wb[_DOMESTIC_SHEET])
        return self._parse_overseas(wb[_OVERSEAS_SHEET])

    def _rows_as_dicts(self, ws) -> list[dict]:
        rows = list(ws.iter_rows(values_only=True))
        header, *data_rows = rows
        return [dict(zip(header, row)) for row in data_rows if any(v is not None for v in row)]

    def _parse_domestic(self, ws) -> list[ParsedTransaction]:
        results = []
        for row in self._rows_as_dicts(ws):
            results.append(
                ParsedTransaction(
                    transaction_date=parse_date(f"{row['승인일자']} {row['승인시각']}"),
                    transaction_time=parse_time(row["승인시각"]),
                    transaction_type="EXPENSE",
                    amount=parse_amount(row["승인금액(원)"]),
                    currency="KRW",
                    merchant_raw=str(row["가맹점명"]).strip(),
                    card_institution="삼성카드",
                    card_number_masked=mask_card_number_if_needed(str(row["카드번호"])),
                    source_transaction_id=str(row["승인번호"]),
                    is_excluded=is_cancelled(row.get("취소여부")),
                    raw_row=row,
                )
            )
        return results

    def _parse_overseas(self, ws) -> list[ParsedTransaction]:
        results = []
        for row in self._rows_as_dicts(ws):
            local_amount = row.get("현지이용금액")
            local_currency = row.get("현지거래통화")
            description = (
                f"{row['업종']} / 현지금액 {local_amount} {local_currency}"
                if local_amount is not None
                else row.get("업종")
            )
            results.append(
                ParsedTransaction(
                    transaction_date=parse_date(f"{row['승인일자']} {row['승인시각']}"),
                    transaction_time=parse_time(row["승인시각"]),
                    transaction_type="EXPENSE",
                    amount=parse_amount(row["승인금액(USD)"]),
                    currency="USD",
                    merchant_raw=str(row["가맹점명"]).strip(),
                    card_institution="삼성카드",
                    card_number_masked=mask_card_number_if_needed(str(row["카드번호"])),
                    source_transaction_id=str(row["승인번호"]),
                    description=description,
                    is_excluded=is_cancelled(row.get("취소구분")),
                    raw_row=row,
                )
            )
        return results
