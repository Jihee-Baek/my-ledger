"""김포페이(김포지역화폐, 코나아이) '거래내역서' Excel export.

Sheet layout: a preamble (title, 조회기간, 성명...), a header row starting
with '거래일시', one row per 충전/결제/결제 취소, and a '...월분 소계' subtotal
row after each month that must be skipped.

Money model - the bank side already records the 충전 as a TRANSFER (the
cash leaves the bank account), so here:

  * 결제       -> EXPENSE for 거래금액 (the price the merchant received)
  * 결제 취소   -> same 승인번호 as the original, is_excluded; the shared id
                  lets mark_cancelled_pairs() exclude the original too
  * 충전       -> TRANSFER for 총 결제금액 (cash actually paid) plus an
                  INCOME row for the bonus (거래금액 - 총 결제금액, the 8~10%
                  할인형 인센티브 credited on top of the charge)
  * per-payment portions funded by 기후행동기회소득 / 걷기 챌린지 columns are
    government incentives that never appear as their own rows, so they are
    recorded as INCOME on the day they were spent (the only time we see them)

The 할인형 인센티브 portion consumed at payment is *not* income again - it
was already counted when the charge bonus was credited.
"""

from decimal import Decimal
from pathlib import Path

import openpyxl

from app.importers.base import ParsedTransaction
from app.importers.normalize import parse_amount, parse_date, parse_time
from app.importers.sniff import is_real_xlsx

_INSTITUTION = "김포페이"
_HEADER_FIRST = "거래일시"
# header cells are truncated in the export ('김포페이 할인형 인센티...'), so match by prefix
_PROGRAM_INCENTIVE_PREFIXES = ("2025년 기후행동", "2026년 기후행동", "김포시 걷기")
_DISCOUNT_INCENTIVE_PREFIX = "김포페이 할인형"


def _num(v) -> Decimal:
    if v in (None, "", "-"):
        return Decimal(0)
    return parse_amount(v)


class GimpoPayImporter:
    source = "gimpo_pay"

    def can_handle(self, file_path: Path) -> bool:
        if not is_real_xlsx(file_path):
            return False
        try:
            ws = openpyxl.load_workbook(file_path, read_only=True).active
            head = [str(r[0] or "") for r in ws.iter_rows(min_row=1, max_row=12, max_col=1, values_only=True)]
        except Exception:
            return False
        return any("김포지역화폐" in h for h in head) and any(h.strip() == _HEADER_FIRST for h in head)

    def parse(self, file_path: Path) -> list[ParsedTransaction]:
        ws = openpyxl.load_workbook(file_path, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))

        header_idx = next(i for i, r in enumerate(rows) if str(r[0] or "").strip() == _HEADER_FIRST)
        header = [str(h or "").replace("\r\n", "").strip() for h in rows[header_idx]]
        col = {name: i for i, name in enumerate(header)}
        program_cols = [i for i, h in enumerate(header) if h.startswith(_PROGRAM_INCENTIVE_PREFIXES)]
        discount_col = next((i for i, h in enumerate(header) if h.startswith(_DISCOUNT_INCENTIVE_PREFIX)), None)

        results: list[ParsedTransaction] = []
        for r in rows[header_idx + 1:]:
            first = str(r[0] or "").strip()
            if len(first) < 10 or not first[:4].isdigit() or first[4] != "/":
                continue  # 소계 / footer rows

            when = first.replace("/", ".")
            tx_date = parse_date(when[:10])
            tx_time = parse_time(when)
            kind = str(r[col["거래방식"]] or "").strip()
            approval = str(r[col["승인번호"]] or "").strip()
            face = _num(r[col["거래금액"]])  # 거래금액
            paid = _num(r[col["총 결제금액"]])
            merchant = str(r[col["가맹점명"]] or "").strip()
            card = str(r[col["카드번호"]] or "").strip()
            program_amounts = {header[i]: _num(r[i]) for i in program_cols if _num(r[i]) != 0}
            raw = {
                "거래일시": first,
                "거래방식": kind,
                "승인번호": approval,
                "거래금액": str(face),
                "총 결제금액": str(paid),
                "충전잔액": str(_num(r[col["충전잔액"]])),
                "할인형 인센티브": str(_num(r[discount_col])) if discount_col is not None else None,
                "지원금": {k: str(v) for k, v in program_amounts.items()},
                "가맹점명": merchant,
            }
            common = dict(
                transaction_date=tx_date,
                transaction_time=tx_time,
                currency="KRW",
                card_institution=_INSTITUTION,
                card_number_masked=card,
                card_type="PREPAID",
                raw_row=raw,
            )

            if kind == "충전":
                results.append(
                    ParsedTransaction(
                        **common,
                        transaction_type="TRANSFER",
                        amount=paid,
                        merchant_raw="김포페이 충전",
                        source_transaction_id=approval,
                        description=f"{merchant} - {face:,.0f}원 충전에 {paid:,.0f}원 결제 (은행 출금과 대응)",
                    )
                )
                bonus = face - paid
                if bonus > 0:
                    results.append(
                        ParsedTransaction(
                            **common,
                            transaction_type="INCOME",
                            amount=bonus,
                            merchant_raw="김포페이 충전 인센티브",
                            source_transaction_id=f"{approval}-bonus",
                            description=f"충전 {face:,.0f}원 중 할인형 인센티브 {bonus:,.0f}원",
                        )
                    )
                continue

            if kind in ("결제", "결제 취소"):
                cancelled = kind == "결제 취소"
                results.append(
                    ParsedTransaction(
                        **common,
                        transaction_type="EXPENSE",
                        amount=face,  # 취소는 음수 그대로 - 원거래와 지문이 겹치지 않아야 두 행 모두 저장된다
                        merchant_raw=merchant,
                        source_transaction_id=approval,
                        description="결제 취소" if cancelled else None,
                        is_excluded=cancelled,
                    )
                )
                for label, amount in program_amounts.items():
                    if cancelled or amount <= 0:
                        continue
                    short = label.split("(")[0].replace("...", "").strip()
                    results.append(
                        ParsedTransaction(
                            **common,
                            transaction_type="INCOME",
                            amount=amount,
                            merchant_raw="김포페이 지원금",
                            source_transaction_id=f"{approval}-{abs(hash(label)) % 10_000}",
                            # 가맹점명은 넣지 않는다 - description도 규칙 매칭 대상이라 '학원' 같은
                            # 가맹점 키워드에 걸려 지원금 수입이 지출 카테고리로 분류된다 (raw_row에 있음)
                            description=f"{short} 지원금으로 결제된 금액",
                        )
                    )
                continue

            raise ValueError(f"김포페이: 알 수 없는 거래방식 {kind!r} ({first})")
        return results
